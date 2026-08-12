"""
'월간 현황 보고서' 양식을 채워서 .docx 파일로 생성하는 모듈.

원본 양식 구조:
  제목: 월간 현황 보고서(YYYY년 M월)         작성자 : 이름
  ┌───────────────────────────────┬─────────────────┐
  │        금월 계획 및 실적          │     차월 계획      │
  ├───────────────┬───────────────┼────────┬────────┤
  │      계획       │      실적       │   계획   │   비고   │
  ├───────────────┼───────────────┼────────┼────────┤
  │   (내용)        │   (내용)        │ (내용)  │  (내용) │
  ├───────────────┴───────────────┴────────┴────────┤
  │                     애로 사항                       │
  ├─────────────────────────────────────────────────┤
  │                   (내용)                          │
  ├─────────────────────────────────────────────────┤
  │                  요청 및 전달 사항                    │
  ├─────────────────────────────────────────────────┤
  │                   (내용)                          │
  └─────────────────────────────────────────────────┘
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_cell_shading(cell, color_hex):
    """셀 배경색을 지정한다 (연두색 헤더 재현용)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _set_cell_text(cell, text, bold=False, align_center=False, font_size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align_center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = paragraph if i == 0 else cell.add_paragraph()
        if align_center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(font_size)


def generate_report_docx(
    output_path,
    year,
    month,
    author,
    plan_this_month,
    result_this_month,
    plan_next_month,
    remark,
    issues,
    requests_text,
):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    # 제목
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(f"월간 현황 보고서({year}년 {month}월)")
    title_run.bold = True
    title_run.font.size = Pt(20)

    # 작성자
    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    author_run = author_p.add_run(f"작성자 : {author}")
    author_run.font.size = Pt(11)

    doc.add_paragraph()

    HEADER_FILL = "E2EFDA"  # 연두색 계열

    # ---- 표 1: 금월 계획/실적 + 차월 계획 ----
    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    col_width = Cm(4.5)
    for row in table.rows:
        for cell in row.cells:
            cell.width = col_width

    # 1행: 대제목 병합 (0~1 병합, 2~3 병합)
    cell_a = table.cell(0, 0).merge(table.cell(0, 1))
    _set_cell_text(cell_a, "금월 계획 및 실적", bold=True, align_center=True)
    _set_cell_shading(cell_a, HEADER_FILL)

    cell_b = table.cell(0, 2).merge(table.cell(0, 3))
    _set_cell_text(cell_b, "차월 계획", bold=True, align_center=True)
    _set_cell_shading(cell_b, HEADER_FILL)

    # 2행: 소제목
    headers = ["계 획", "실 적", "계 획", "비 고"]
    for i, h in enumerate(headers):
        cell = table.cell(1, i)
        _set_cell_text(cell, h, bold=True, align_center=True)
        _set_cell_shading(cell, HEADER_FILL)

    # 3행: 내용
    contents = [plan_this_month, result_this_month, plan_next_month, remark]
    for i, c in enumerate(contents):
        _set_cell_text(table.cell(2, i), c, font_size=10)

    doc.add_paragraph()

    # ---- 표 2: 애로사항 / 요청 및 전달사항 ----
    table2 = doc.add_table(rows=4, cols=1)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.style = "Table Grid"
    for row in table2.rows:
        row.cells[0].width = Cm(18.0)

    _set_cell_text(table2.cell(0, 0), "애로 사항", bold=True, align_center=True)
    _set_cell_shading(table2.cell(0, 0), HEADER_FILL)
    _set_cell_text(table2.cell(1, 0), issues, font_size=10)

    _set_cell_text(table2.cell(2, 0), "요청 및 전달 사항", bold=True, align_center=True)
    _set_cell_shading(table2.cell(2, 0), HEADER_FILL)
    _set_cell_text(table2.cell(3, 0), requests_text, font_size=10)

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    # 간단한 동작 테스트
    generate_report_docx(
        "test_report.docx",
        year=2026,
        month=8,
        author="홍길동",
        plan_this_month="[업무]\n1. 테스트 프로젝트\n- 항목 1\n- 항목 2",
        result_this_month="[업무]\n1. 테스트 프로젝트\n- 완료된 항목 1",
        plan_next_month="[업무]\n1. 다음달 계획 항목",
        remark="",
        issues="특이사항 없음",
        requests_text="없음",
    )
    print("test_report.docx 생성 완료")
